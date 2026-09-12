"""Telco analysis: correlations, visuals, logistic churn model, PDF report.
Usage: python scripts/analysis_telco.py
"""
import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from matplotlib.backends.backend_pdf import PdfPages
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score, precision_score, recall_score, roc_auc_score, roc_curve, confusion_matrix
import joblib

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DATA_FILE = os.path.join(ROOT, 'data', 'processed', 'telco_cleaned.csv')
REPORTS_DIR = os.path.join(ROOT, 'reports')
MODELS_DIR = os.path.join(ROOT, 'models')
os.makedirs(REPORTS_DIR, exist_ok=True)
os.makedirs(MODELS_DIR, exist_ok=True)


def load_data():
    df = pd.read_csv(DATA_FILE)
    return df


def numeric_correlation(df):
    num = df.select_dtypes(include=[np.number]).copy()
    # drop id-like columns
    for c in ['Count']:
        if c in num.columns:
            num = num.drop(columns=[c])
    corr = num.corr()
    plt.figure(figsize=(10,8))
    sns.heatmap(corr, annot=False, cmap='coolwarm', center=0)
    fn = os.path.join(REPORTS_DIR, 'correlation_heatmap.png')
    plt.title('Numeric Feature Correlation')
    plt.tight_layout()
    plt.savefig(fn)
    plt.close()
    return fn, corr


def create_visuals(df):
    files = []
    # Churn count
    plt.figure(figsize=(6,4))
    sns.countplot(x='Churn', data=df)
    plt.title('Churn counts')
    f = os.path.join(REPORTS_DIR, 'churn_counts.png')
    plt.tight_layout(); plt.savefig(f); plt.close(); files.append(f)

    # Monthly Charges distribution (column may be 'Monthly Charges')
    mc_col = None
    for c in df.columns:
        if c.replace(' ', '').lower() == 'monthlycharges' or c.lower() == 'monthly charges':
            mc_col = c
            break
    if mc_col is not None:
        plt.figure(figsize=(6,4))
        sns.kdeplot(data=df, x=mc_col, hue='Churn', common_norm=False)
        plt.title('Monthly Charges by Churn')
        f = os.path.join(REPORTS_DIR, 'monthlycharges_dist.png')
    plt.tight_layout(); plt.savefig(f); plt.close(); files.append(f)

    # Tenure histogram (handle 'Tenure Months' or similar)
    tenure_col = next((c for c in df.columns if 'tenure' in c.lower()), None)
    if tenure_col is not None:
        plt.figure(figsize=(6,4))
        sns.histplot(df[tenure_col].dropna(), bins=30)
        plt.title('Tenure distribution')
        f = os.path.join(REPORTS_DIR, 'tenure_hist.png')
        plt.tight_layout(); plt.savefig(f); plt.close(); files.append(f)

    # Churn rate by Contract
    if 'Contract' in df.columns:
        rates = df.groupby('Contract')['ChurnBinary'].mean().sort_values(ascending=False)
        plt.figure(figsize=(6,4))
        rates.plot(kind='bar')
        plt.ylabel('Churn rate')
        plt.title('Churn rate by Contract')
        f = os.path.join(REPORTS_DIR, 'churn_by_contract.png')
        plt.tight_layout(); plt.savefig(f); plt.close(); files.append(f)

    return files


def train_model(df):
    # Prepare features: drop identifiers, leak columns, target
    X = df.copy()
    if 'CustomerID' in X.columns:
        X = X.drop(columns=['CustomerID'])
    # Drop any column that contains 'churn' (case-insensitive) except the target 'ChurnBinary'
    churn_like = [c for c in X.columns if 'churn' in str(c).lower() and c != 'ChurnBinary']
    if churn_like:
        X = X.drop(columns=churn_like)
    y = X.pop('ChurnBinary')

    # Select numeric features and simple encoding for categoricals
    num_cols = X.select_dtypes(include=[np.number]).columns.tolist()
    cat_cols = X.select_dtypes(include=['category', 'object']).columns.tolist()
    # Remove target-looking numeric cols
    # remove raw churn label if present
    if any(str(c).strip().lower() == 'churn' for c in num_cols):
        num_cols = [c for c in num_cols if str(c).strip().lower() != 'churn']

    X_num = X[num_cols].fillna(0)
    X_cat = pd.get_dummies(X[cat_cols].fillna('NA'), drop_first=True)
    X_all = pd.concat([X_num, X_cat], axis=1)

    # Train/test
    X_train, X_test, y_train, y_test = train_test_split(X_all, y, test_size=0.2, random_state=42, stratify=y)

    scaler = StandardScaler()
    X_train_scaled = scaler.fit_transform(X_train)
    X_test_scaled = scaler.transform(X_test)

    model = LogisticRegression(max_iter=1000)
    model.fit(X_train_scaled, y_train)

    # Predictions and metrics
    y_pred = model.predict(X_test_scaled)
    y_proba = model.predict_proba(X_test_scaled)[:,1]
    metrics = {
        'accuracy': accuracy_score(y_test, y_pred),
        'precision': precision_score(y_test, y_pred, zero_division=0),
        'recall': recall_score(y_test, y_pred, zero_division=0),
        'roc_auc': roc_auc_score(y_test, y_proba)
    }

    # ROC plot
    fpr, tpr, _ = roc_curve(y_test, y_proba)
    plt.figure(figsize=(6,4))
    plt.plot(fpr, tpr, label=f"ROC AUC={metrics['roc_auc']:.3f}")
    plt.plot([0,1],[0,1],'k--')
    plt.xlabel('FPR'); plt.ylabel('TPR'); plt.title('ROC curve'); plt.legend()
    roc_fn = os.path.join(REPORTS_DIR, 'roc_curve.png')
    plt.tight_layout(); plt.savefig(roc_fn); plt.close()

    # Save model and scaler
    model_path = os.path.join(MODELS_DIR, 'logistic_telco.pkl')
    joblib.dump({'model': model, 'scaler': scaler, 'features': X_all.columns.tolist()}, model_path)

    return metrics, roc_fn, model_path


def make_pdf_report(fig_files, corr_fig, metrics, df):
    pdf_path = os.path.join(REPORTS_DIR, 'telco_report.pdf')
    with PdfPages(pdf_path) as pdf:
        # Title page
        plt.figure(figsize=(8.27, 11.69))
        plt.axis('off')
        plt.text(0.5, 0.6, 'Telco Customer Churn — EDA & Model', ha='center', fontsize=20)
        plt.text(0.5, 0.55, f"Rows: {df.shape[0]}  |  Churn rate: {df['ChurnBinary'].mean():.3f}", ha='center')
        pdf.savefig(); plt.close()

        # Add correlation
        corr_img = plt.imread(corr_fig)
        plt.figure(figsize=(8,6)); plt.imshow(corr_img); plt.axis('off'); pdf.savefig(); plt.close()

        # Add each figure
        for f in fig_files:
            img = plt.imread(f)
            plt.figure(figsize=(8,6)); plt.imshow(img); plt.axis('off'); pdf.savefig(); plt.close()

        # Metrics page
        plt.figure(figsize=(8.27,11.69)); plt.axis('off')
        txt = 'Model metrics:\n' + '\n'.join([f"{k}: {v:.3f}" for k,v in metrics.items()])
        plt.text(0.1, 0.8, txt, fontsize=12)
        pdf.savefig(); plt.close()

    return pdf_path


if __name__ == '__main__':
    print('Loading data...')
    df = load_data()
    print('Creating visuals...')
    corr_fig, corr = numeric_correlation(df)
    figs = create_visuals(df)
    print('Training model...')
    metrics, roc_fn, model_path = train_model(df)
    print('Model metrics:', metrics)
    print('Creating PDF report...')
    pdf = make_pdf_report(figs, corr_fig, metrics, df)
    print('Saved report to', pdf)
    print('Saved model to', model_path)
    # Attempt simple PPTX (optional)
    try:
        from pptx import Presentation
        prs = Presentation()
        sld = prs.slides.add_slide(prs.slide_layouts[5])
        title = sld.shapes.title
        title.text = 'Telco Churn — Key Findings'
        tx = sld.shapes.add_textbox(left=1000000, top=1500000, width=6000000, height=2000000)
        tf = tx.text_frame
        tf.text = f"Rows: {df.shape[0]}  Churn rate: {df['ChurnBinary'].mean():.3f}"
        pptx_path = os.path.join(REPORTS_DIR, 'telco_report.pptx')
        prs.save(pptx_path)
        print('Saved PPTX to', pptx_path)
    except Exception:
        print('python-pptx not available — skipped PPTX generation')
